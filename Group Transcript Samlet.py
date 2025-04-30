import json
from docx import Document

def group_segments_by_speaker(json_file_path, output_docx_path, max_pause_seconds=4):
    with open(json_file_path, "r", encoding="utf-8") as f:
        segments = json.load(f)["segments"]

    grouped_segments = []
    current_speaker = None
    start_time = None
    end_time = None
    text_buffer = []
    last_end = 0

    for seg in segments:
        speaker = seg.get("speaker", "Undetermined")
        text = seg["text"].strip()
        start = seg["start"]
        end = seg["end"]

        if speaker == "Undetermined" or text == "...":
            continue

        if speaker == current_speaker and (start - last_end) <= max_pause_seconds:
            text_buffer.append(text)
            end_time = end
        else:
            if current_speaker and text_buffer:
                grouped_segments.append((start_time, end_time, current_speaker, " ".join(text_buffer)))
            current_speaker = speaker
            start_time = start
            end_time = end
            text_buffer = [text]
        last_end = end

    if current_speaker and text_buffer:
        grouped_segments.append((start_time, end_time, current_speaker, " ".join(text_buffer)))

    doc = Document()
    doc.add_heading("Interview – Grupperet og Tidskodet", 0)

    for start, end, speaker, text in grouped_segments:
        t_start = f"{int(start // 60):02}:{int(start % 60):02}"
        t_end = f"{int(end // 60):02}:{int(end % 60):02}"
        para = doc.add_paragraph()
        para.add_run(f"{t_start} - {t_end} {speaker}:
").bold = True
        para.add_run(text)

    doc.save(output_docx_path)